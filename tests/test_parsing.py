"""CSV-in-ZIP parsing."""
from __future__ import annotations

import io
import zipfile

import pandas as pd
import pytest

from wgea_mcp.parsing import (
    ParseError,
    drop_blank_rows,
    list_zip_members,
    read_csv_from_zip,
    stream_csv_from_zip,
)


def _make_zip(members: dict[str, str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, content in members.items():
            zf.writestr(name, content)
    return buf.getvalue()


def test_list_zip_members(sample_zip_bytes):
    members = list_zip_members(sample_zip_bytes)
    assert any("workforce_composition" in m for m in members)
    assert any("harm_prevention" in m for m in members)


def test_list_zip_members_empty_body():
    with pytest.raises(ParseError, match="empty ZIP"):
        list_zip_members(b"")


def test_list_zip_members_corrupt_body():
    with pytest.raises(ParseError, match="corrupt"):
        list_zip_members(b"not a zip file at all")


def test_read_csv_exact_name(sample_zip_bytes):
    df = read_csv_from_zip(sample_zip_bytes, "wgea_workforce_composition_2025.csv")
    assert isinstance(df, pd.DataFrame)
    assert len(df) > 0
    assert "employer_name" in df.columns
    assert "n_employees" in df.columns


def test_read_csv_prefix_match(sample_zip_bytes):
    """`zip_member: wgea_workforce_composition_` should match the year-suffixed file."""
    df = read_csv_from_zip(sample_zip_bytes, "wgea_workforce_composition_")
    assert "employer_name" in df.columns


def test_read_csv_max_rows(sample_zip_bytes):
    df = read_csv_from_zip(sample_zip_bytes, "wgea_workforce_composition_", max_rows=5)
    assert len(df) == 5


def test_read_csv_missing_member(sample_zip_bytes):
    with pytest.raises(ParseError, match="no member in ZIP matches"):
        read_csv_from_zip(sample_zip_bytes, "wgea_does_not_exist")


def test_read_csv_bom_stripped():
    body = _make_zip({"x.csv": "﻿a,b,c\n1,2,3\n"})
    df = read_csv_from_zip(body, "x.csv")
    assert list(df.columns) == ["a", "b", "c"]


def test_drop_blank_rows():
    df = pd.DataFrame({"a": [1, None, 3], "b": ["x", None, "z"]})
    out = drop_blank_rows(df, ["a", "b"])
    assert len(out) == 2
    assert list(out["a"]) == [1, 3]


def test_drop_blank_rows_no_matching_columns():
    df = pd.DataFrame({"a": [1, None]})
    out = drop_blank_rows(df, ["b"])
    # No matching cols → return unchanged
    assert len(out) == 2


def test_read_csv_year_disambiguation():
    """When multiple year-suffixed members match a prefix, pick the newest."""
    body = _make_zip({
        "wgea_x_2024.csv": "a\n1\n",
        "wgea_x_2025.csv": "a\n2\n",
        "wgea_x_2023.csv": "a\n3\n",
    })
    df = read_csv_from_zip(body, "wgea_x_")
    assert df["a"].iloc[0] == 2  # newest year wins


def test_read_csv_regex_fallback():
    body = _make_zip({"foo_bar_baz.csv": "a\n1\n"})
    df = read_csv_from_zip(body, r"foo.*baz")
    assert df["a"].iloc[0] == 1


def test_stream_csv_short_circuits_on_max_rows():
    """stream_csv_from_zip must stop reading once max_rows is hit."""
    rows = "\n".join(f"{i},x" for i in range(1000))
    body = _make_zip({"x.csv": f"a,b\n{rows}\n"})
    df = stream_csv_from_zip(body, "x.csv", max_rows=5)
    assert len(df) == 5
    assert list(df["a"]) == [0, 1, 2, 3, 4]


def test_stream_csv_applies_row_predicate():
    body = _make_zip({"x.csv": "a,b\n1,red\n2,blue\n3,red\n4,blue\n5,red\n"})
    df = stream_csv_from_zip(
        body, "x.csv", max_rows=10, row_predicate=lambda r: r["b"] == "red"
    )
    assert list(df["a"]) == [1, 3, 5]


def test_stream_csv_short_circuits_with_predicate():
    """When max_rows of MATCHING rows are seen, stop iterating even mid-file."""
    body = _make_zip({"x.csv": "a,b\n1,red\n2,blue\n3,red\n4,red\n5,red\n"})
    seen = {"count": 0}

    def pred(r):
        seen["count"] += 1
        return r["b"] == "red"

    df = stream_csv_from_zip(body, "x.csv", max_rows=2, row_predicate=pred)
    assert len(df) == 2
    # Should have read rows 1, 2, 3 (1=red, 2=blue, 3=red — stops on 2nd red).
    assert seen["count"] == 3


def test_stream_csv_max_rows_required():
    body = _make_zip({"x.csv": "a\n1\n"})
    with pytest.raises(ValueError, match="max_rows"):
        stream_csv_from_zip(body, "x.csv", max_rows=0)


def test_stream_csv_empty_result_returns_typed_df():
    """Predicate that rejects every row should return an empty DataFrame
    with the source columns preserved."""
    body = _make_zip({"x.csv": "a,b\n1,red\n2,blue\n"})
    df = stream_csv_from_zip(
        body, "x.csv", max_rows=10, row_predicate=lambda r: False
    )
    assert len(df) == 0
    assert list(df.columns) == ["a", "b"]


def test_stream_csv_numeric_inference():
    """Numeric-looking columns should coerce to numeric dtype, mirroring
    the full-parse path's pandas inference."""
    body = _make_zip({"x.csv": "a,b\n1,foo\n2,bar\n3,baz\n"})
    df = stream_csv_from_zip(body, "x.csv", max_rows=10)
    assert pd.api.types.is_numeric_dtype(df["a"])
    # String column stays string.
    assert df["b"].iloc[0] == "foo"


# ─── HEADLINE_GAP: parse_egpg_xlsx aggregator tests ─────────────────────

from wgea_mcp.parsing import parse_egpg_xlsx


def test_egpg_xlsx_extracts_reporting_year(sample_egpg_xlsx_bytes):
    _df, year = parse_egpg_xlsx(sample_egpg_xlsx_bytes)
    assert year == "2024-25"


def test_egpg_xlsx_returns_industry_rows_plus_all(sample_egpg_xlsx_bytes):
    """Fixture has 3 industries (Mining, Finance, Construction) + 1 Commonwealth
    row that must be filtered out. Aggregator yields 3 industry rows + 1
    synthetic 'All employers' row = 4 rows."""
    df, _year = parse_egpg_xlsx(sample_egpg_xlsx_bytes)
    assert len(df) == 4
    divisions = set(df["anzsic_division"])
    assert divisions == {
        "All employers",
        "Mining",
        "Financial and Insurance Services",
        "Construction",
    }
    # Commonwealth public-sector row must NOT leak through.
    assert "Public Administration and Safety" not in divisions


def test_egpg_xlsx_all_employers_row_is_first(sample_egpg_xlsx_bytes):
    """HEADLINE_GAP latest() with no filter should surface the national
    'All employers' row first."""
    df, _year = parse_egpg_xlsx(sample_egpg_xlsx_bytes)
    assert df.iloc[0]["anzsic_division"] == "All employers"


def test_egpg_xlsx_medians_match_expected(sample_egpg_xlsx_bytes):
    """Fixture's Mining rows have avg_total_rem fractions 0.20/0.15/0.10.
    Median is 0.15 → 15.0%. Confirm the aggregator multiplies by 100."""
    df, _year = parse_egpg_xlsx(sample_egpg_xlsx_bytes)
    mining = df[df["anzsic_division"] == "Mining"].iloc[0]
    assert mining["total_remuneration_gap_pct"] == 15.0
    # base salary: 0.18/0.13/0.08 → median 0.13 → 13.0%
    assert mining["base_salary_gap_pct"] == 13.0
    # median total rem: 0.22/0.18/0.12 → median 0.18 → 18.0%
    assert mining["median_total_rem_gap_pct"] == 18.0


def test_egpg_xlsx_employer_count_is_int(sample_egpg_xlsx_bytes):
    df, _year = parse_egpg_xlsx(sample_egpg_xlsx_bytes)
    for _, row in df.iterrows():
        assert isinstance(row["employer_count"], int), row["anzsic_division"]
        assert row["employer_count"] >= 1, row["anzsic_division"]


def test_egpg_xlsx_all_employers_count_sums_all_industries(sample_egpg_xlsx_bytes):
    """The 'All employers' row's employer_count equals the sum of all
    industry-row employer counts (every private-sector employer is counted
    in both its industry and the national rollup)."""
    df, _year = parse_egpg_xlsx(sample_egpg_xlsx_bytes)
    all_emp = df[df["anzsic_division"] == "All employers"].iloc[0]
    industry_sum = df[df["anzsic_division"] != "All employers"]["employer_count"].sum()
    assert all_emp["employer_count"] == industry_sum


def test_egpg_xlsx_empty_body_raises():
    from wgea_mcp.parsing import ParseError
    with pytest.raises(ParseError, match="empty xlsx body"):
        parse_egpg_xlsx(b"")


def test_egpg_xlsx_corrupt_body_raises():
    from wgea_mcp.parsing import ParseError
    with pytest.raises(ParseError):
        parse_egpg_xlsx(b"not a real xlsx - just some bytes")


def test_egpg_xlsx_missing_column_raises(tmp_path):
    """If WGEA renames a column we want a loud ParseError, not a silent drop."""
    import openpyxl
    from wgea_mcp.parsing import ParseError
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2. Employers "
    ws.cell(row=2, column=1, value="EMPLOYER GENDER PAY GAPS")
    ws.cell(row=3, column=1, value="Results based on 2024-25 WGEA Gender Equality Reporting")
    # Headers DELIBERATELY missing 'Average total remuneration GPG (%)'
    headers = [
        "Employer name", "Employer ABN", "Sector",
        "Industry (ANZSIC Division)", "Average base salary GPG (%)",
        "Median total remuneration GPG (%)", "Median base salary GPG (%)",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=h)
    ws.cell(row=5, column=1, value="Some Co")
    ws.cell(row=5, column=2, value="12345678901")
    ws.cell(row=5, column=3, value="Private")
    ws.cell(row=5, column=4, value="Mining")
    p = tmp_path / "renamed.xlsx"
    wb.save(p)
    with pytest.raises(ParseError, match="missing expected columns"):
        parse_egpg_xlsx(p.read_bytes())


def test_egpg_xlsx_no_private_sector_raises(tmp_path):
    """A spreadsheet with only Commonwealth rows can't yield a meaningful
    mid-point — must raise rather than emit an empty 'All employers' row."""
    import openpyxl
    from wgea_mcp.parsing import ParseError
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2. Employers "
    ws.cell(row=2, column=1, value="EMPLOYER GENDER PAY GAPS")
    ws.cell(row=3, column=1, value="Results based on 2024-25 WGEA Gender Equality Reporting")
    headers = [
        "Employer name", "Employer ABN", "Sector",
        "Industry (ANZSIC Division)", "Industry (ANZSIC Class)",
        "Employer size range   (# employees)",
        "Average total remuneration GPG (%)", "Average base salary GPG (%)",
        "Median total remuneration GPG (%)", "Median base salary GPG (%)",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=h)
    ws.cell(row=5, column=1, value="Public Co")
    ws.cell(row=5, column=2, value="44444444444")
    ws.cell(row=5, column=3, value="Public")
    ws.cell(row=5, column=4, value="Public Administration and Safety")
    ws.cell(row=5, column=5, value="Public")
    ws.cell(row=5, column=6, value="5000+")
    for c in (7, 8, 9, 10):
        ws.cell(row=5, column=c, value=0.05)
    p = tmp_path / "public_only.xlsx"
    wb.save(p)
    with pytest.raises(ParseError, match="Private"):
        parse_egpg_xlsx(p.read_bytes())
